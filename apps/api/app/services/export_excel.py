import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.document import Document
from app.models.project import Project
from app.models.requirement import Requirement

_STATUS_FILL = {
    "compliant_looking": "C6EFCE",
    "human_verified": "C6EFCE",
    "partially_addressed": "FFEB9C",
    "not_applicable_pending_verification": "FFEB9C",
    "evidence_not_found": "FFC7CE",
    "potential_non_compliance": "FFC7CE",
    "human_rejected": "FFC7CE",
    "not_assessed": "D9D9D9",
}

_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

MATRIX_HEADERS = [
    "Clause",
    "Requirement",
    "Category",
    "Mandatory Status",
    "Source Page",
    "Assessment Status",
    "Priority",
    "Evidence Quality",
    "Confidence",
    "AI Explanation",
    "Missing Information",
    "Evidence Excerpts (bid doc, page)",
    "Requires Human Review",
    "Reviewer Status",
]


async def _load_requirements(db: AsyncSession, *, project_id: str) -> list[Requirement]:
    result = await db.scalars(
        select(Requirement)
        .where(Requirement.project_id == project_id)
        .options(
            selectinload(Requirement.assessment),
            selectinload(Requirement.evidence_records),
        )
        .order_by(Requirement.source_page, Requirement.created_at)
    )
    return list(result)


async def build_compliance_matrix_workbook(db: AsyncSession, *, project: Project) -> bytes:
    requirements = await _load_requirements(db, project_id=project.id)
    documents = list(
        await db.scalars(select(Document).where(Document.project_id == project.id))
    )

    wb = Workbook()

    _build_summary_sheet(wb, project=project, requirements=requirements, documents=documents)
    _build_matrix_sheet(wb, requirements=requirements)
    _build_assumptions_sheet(wb, project=project)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _build_summary_sheet(wb: Workbook, *, project: Project, requirements, documents) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws.freeze_panes = "A2"

    rows: list[tuple[str, object]] = [
        ("Project", project.name),
        ("Client", project.client_name or "—"),
        ("Tender reference", project.tender_reference or "—"),
        ("Sector", project.sector or "—"),
        ("Status", project.status.value if hasattr(project.status, "value") else project.status),
        ("Sample / demo data", "Yes — fictional, for demonstration only" if project.is_demo else "No"),
        ("", ""),
        ("Documents", ""),
    ]
    for doc in documents:
        doc_label = f"  {doc.document_type.value}"
        doc_value = f"{doc.original_filename} ({doc.page_count or '?'} pages)"
        rows.append((doc_label, doc_value))

    rows.append(("", ""))
    rows.append(("Total requirements", len(requirements)))

    status_counts: dict[str, int] = {}
    for r in requirements:
        s = r.assessment.status.value if r.assessment and hasattr(r.assessment.status, "value") else (
            r.assessment.status if r.assessment else "not_assessed"
        )
        status_counts[s] = status_counts.get(s, 0) + 1
    for status, count in sorted(status_counts.items()):
        rows.append((f"  {status.replace('_', ' ').title()}", count))

    for row_idx, (row_label, row_value) in enumerate(rows, start=1):
        is_section_header = bool(row_label) and not row_label.startswith(" ")
        ws.cell(row=row_idx, column=1, value=row_label).font = Font(bold=is_section_header)
        ws.cell(row=row_idx, column=2, value=row_value)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60


def _build_matrix_sheet(wb: Workbook, *, requirements: list[Requirement]) -> None:
    ws = wb.create_sheet("Compliance Matrix")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(MATRIX_HEADERS))}1"

    for col_idx, header in enumerate(MATRIX_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_idx, req in enumerate(requirements, start=2):
        assessment = req.assessment
        status = (
            assessment.status.value if assessment and hasattr(assessment.status, "value") else
            (assessment.status if assessment else "not_assessed")
        )
        priority = (
            assessment.priority.value if assessment and hasattr(assessment.priority, "value") else
            (assessment.priority if assessment else "")
        )
        evidence_quality = (
            assessment.evidence_quality.value
            if assessment and hasattr(assessment.evidence_quality, "value")
            else (assessment.evidence_quality if assessment else "")
        )
        confidence = (
            assessment.assessment_confidence.value
            if assessment and hasattr(assessment.assessment_confidence, "value")
            else (assessment.assessment_confidence if assessment else "")
        )
        missing_info = []
        if assessment and assessment.missing_information:
            try:
                missing_info = json.loads(assessment.missing_information)
            except (TypeError, ValueError):
                missing_info = []

        evidence_text = "\n".join(
            f"p.{ev.page_number}: {ev.excerpt}" for ev in req.evidence_records
        ) or "Evidence not located in the provided documents."

        values = [
            req.source_clause or "",
            req.title,
            req.category,
            req.mandatory_status.value if hasattr(req.mandatory_status, "value") else req.mandatory_status,
            req.source_page,
            status.replace("_", " ").title(),
            str(priority).title(),
            str(evidence_quality).title(),
            str(confidence).title(),
            assessment.reason if assessment else "",
            "; ".join(missing_info),
            evidence_text,
            "Yes" if (assessment and assessment.requires_human_review) else "No",
            assessment.reviewer_status.title() if assessment else "Unreviewed",
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        fill_color = _STATUS_FILL.get(status)
        if fill_color:
            ws.cell(row=row_idx, column=6).fill = PatternFill(
                start_color=fill_color, end_color=fill_color, fill_type="solid"
            )

    widths = [10, 40, 16, 14, 10, 20, 12, 14, 12, 50, 30, 50, 12, 14]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _build_assumptions_sheet(wb: Workbook, *, project: Project) -> None:
    ws = wb.create_sheet("Assumptions & Limitations")
    ws.column_dimensions["A"].width = 100
    lines = [
        "TenderGuard — Assumptions & Limitations",
        "",
        "This export is generated from AI-assisted analysis and human review recorded in "
        "TenderGuard. It is a decision-support tool, not a compliance guarantee or a "
        "substitute for legal, technical, or commercial review.",
        "",
        "- 'Evidence not located' means no supporting text was found in the documents "
        "provided at the time of analysis — it is not automatic proof of non-compliance.",
        "- AI-generated assessments (Reviewer Status = 'Unreviewed') have not yet been "
        "confirmed by a human reviewer and should be verified before submission.",
        "- Priority is derived from mandatory status and assessment status using a "
        "documented, configurable rule (see docs/architecture.md) — it is not a legal or "
        "procurement determination.",
    ]
    if project.is_demo:
        lines += [
            "",
            "This project uses fictional sample data created for demonstration purposes. "
            "The tender, bidder, company names, and all figures are invented and do not "
            "describe a real procurement.",
        ]
    for row_idx, line in enumerate(lines, start=1):
        cell = ws.cell(row=row_idx, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row_idx == 1:
            cell.font = Font(bold=True, size=14)
