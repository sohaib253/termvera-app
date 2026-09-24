"""ClauseRisk Excel export — consolidates the brief's 7 requested reports
into one workbook's worth of sheets (Executive Summary, Clause Risk
Register, Commercial Risk Summary, Liability Exposure Summary,
Negotiation Issues List, Assumptions & Limitations), the same pragmatic
MVP consolidation TenderGuard's app/services/export_excel.py uses rather
than building 7 separate documents.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.contract import Contract, ContractVersion
from app.services.clauserisk import contracts as contract_service

_SEVERITY_FILL = {
    "critical": "FFC7CE",
    "high": "FFEB9C",
    "medium": "FFEB9C",
    "low": "C6EFCE",
    "informational": "D9D9D9",
}
_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

REGISTER_HEADERS = [
    "Clause",
    "Category",
    "Risk Type",
    "Severity",
    "Score",
    "Description",
    "Contractual Effect",
    "Potential Exposure",
    "Affected Party",
    "Uncertainty",
    "Recommended Action",
    "Reviewer Status",
]


async def build_risk_report_workbook(
    db: AsyncSession, *, contract: Contract, version: ContractVersion
) -> bytes:
    findings = await contract_service.list_risk_findings(db, contract_version_id=version.id)

    wb = Workbook()
    _build_executive_summary(wb, contract=contract, version=version, findings=findings)
    _build_register_sheet(wb, "Clause Risk Register", findings)
    _build_register_sheet(
        wb, "Commercial Risk Summary", [f for f in findings if f.category == "commercial"]
    )
    _build_register_sheet(
        wb, "Liability Exposure Summary", [f for f in findings if f.category == "liability"]
    )
    _build_negotiation_sheet(wb, findings)
    _build_assumptions_sheet(wb)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _build_executive_summary(wb: Workbook, *, contract, version, findings) -> None:
    ws = wb.active
    ws.title = "Executive Summary"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 60

    severity_counts: dict[str, int] = {}
    for f in findings:
        severity_counts[f.severity.value] = severity_counts.get(f.severity.value, 0) + 1

    rows: list[tuple[str, object]] = [
        ("Contract", contract.name),
        ("Counterparty", contract.counterparty_name or "—"),
        ("Contract type", contract.contract_type.value),
        ("Version reviewed", version.version_label),
        ("", ""),
        ("Total risk findings", len(findings)),
    ]
    for severity in ("critical", "high", "medium", "low", "informational"):
        rows.append((f"  {severity.title()}", severity_counts.get(severity, 0)))

    for row_idx, (label, value) in enumerate(rows, start=1):
        is_header = bool(label) and not label.startswith(" ")
        ws.cell(row=row_idx, column=1, value=label).font = Font(bold=is_header)
        ws.cell(row=row_idx, column=2, value=value)


def _build_register_sheet(wb: Workbook, title: str, findings: list) -> None:
    ws = wb.create_sheet(title[:31])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(REGISTER_HEADERS))}1"

    for col_idx, header in enumerate(REGISTER_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_idx, finding in enumerate(findings, start=2):
        clause_label = finding.clause.clause_number if finding.clause else "(contract-wide)"
        values = [
            clause_label,
            finding.category,
            finding.risk_type,
            finding.severity.value.title(),
            finding.computed_score,
            finding.risk_description,
            finding.contractual_effect or "",
            finding.potential_exposure or "",
            finding.affected_party.value,
            finding.uncertainty.value.replace("_", " ").title(),
            finding.recommended_review_action or "",
            finding.reviewer_status.value.replace("_", " ").title(),
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        fill_color = _SEVERITY_FILL.get(finding.severity.value)
        if fill_color:
            ws.cell(row=row_idx, column=4).fill = PatternFill(
                start_color=fill_color, end_color=fill_color, fill_type="solid"
            )

    widths = [10, 14, 24, 12, 8, 45, 30, 30, 12, 16, 30, 16]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _build_negotiation_sheet(wb: Workbook, findings: list) -> None:
    negotiation_items = [
        f
        for f in findings
        if f.reviewer_status.value == "marked_for_negotiation"
        or f.severity.value in ("critical", "high")
    ]
    _build_register_sheet(wb, "Negotiation Issues List", negotiation_items)


def _build_assumptions_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Assumptions & Limitations")
    ws.column_dimensions["A"].width = 100
    lines = [
        "ClauseRisk — Assumptions & Limitations",
        "",
        "This export is generated from AI-assisted contract analysis and human review "
        "recorded in ClauseRisk. It is a contract-review decision-support tool, not legal "
        "advice, and does not replace legal or commercial review.",
        "",
        "- Severity and score are computed by a documented, deterministic rule engine "
        "(app/services/clauserisk/risk_engine.py) that adjusts the AI's initial suggestion — "
        "they are not the AI's judgment alone.",
        "- 'Uncertainty: not found' means no relevant clause or evidence was located in the "
        "reviewed documents at the time of analysis — it is not proof the underlying "
        "protection or obligation does not exist elsewhere in the contract.",
        "- Findings with Reviewer Status 'Unreviewed' have not yet been confirmed by a human "
        "reviewer and should be verified before relying on them.",
        "- Clause segmentation is based on detected numbered clause headers and may "
        "under-segment contracts with unconventional structure.",
    ]
    for row_idx, line in enumerate(lines, start=1):
        cell = ws.cell(row=row_idx, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row_idx == 1:
            cell.font = Font(bold=True, size=14)
