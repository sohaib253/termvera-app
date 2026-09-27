"""Extraction of scanned documents: PDFs with no text layer, rotated
scans, and image uploads. The OCR tests need the Tesseract executable and
are skipped where it isn't installed."""

import io

import fitz
import pytest
from PIL import Image, ImageDraw, ImageFont

from app.services.extraction import extract_document
from app.services.ocr import ocr_available

requires_ocr = pytest.mark.skipif(not ocr_available(), reason="Tesseract is not installed")

CLAUSE = "SECTION 15. TERMINATION"
BODY_LINES = [
    "The Company shall have the right to terminate",
    "the Contract upon thirty days notice.",
]


def _scanned_page_image(rotate: int = 0) -> Image.Image:
    """A page of typewritten text as a scanner would produce it: pixels only."""
    image = Image.new("L", (1700, 2200), color=255)
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 36)
    except OSError:
        font = ImageFont.load_default(size=36)
    draw.text((150, 200), CLAUSE, fill=0, font=font)
    for i, line in enumerate(BODY_LINES):
        draw.text((150, 280 + i * 60), line, fill=0, font=font)
    return image.rotate(rotate, expand=True, fillcolor=255) if rotate else image


def _pdf_from_images(*images: Image.Image) -> bytes:
    pdf = fitz.open()
    for image in images:
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        page = pdf.new_page(width=612, height=792)
        page.insert_image(page.rect, stream=buffer.getvalue())
    out = io.BytesIO()
    pdf.save(out)
    return out.getvalue()


def _normalise(text: str) -> str:
    return " ".join(text.split()).lower()


@requires_ocr
def test_scanned_pdf_without_text_layer_is_read_by_ocr():
    result = extract_document(_pdf_from_images(_scanned_page_image()), filename="scan.pdf")

    page = result.pages[0]
    assert page.extraction_method == "ocr"
    assert not page.low_text_warning
    assert "termination" in _normalise(page.text)
    assert "thirty days notice" in _normalise(page.text)


@requires_ocr
def test_upside_down_scan_is_rotated_before_reading():
    result = extract_document(_pdf_from_images(_scanned_page_image(rotate=180)), filename="scan.pdf")

    assert "thirty days notice" in _normalise(result.pages[0].text)


@requires_ocr
def test_image_upload_is_read_by_ocr():
    buffer = io.BytesIO()
    _scanned_page_image().save(buffer, format="PNG")

    result = extract_document(buffer.getvalue(), filename="page.png")

    assert result.page_count == 1
    assert "termination" in _normalise(result.pages[0].text)


@requires_ocr
def test_progress_is_reported_for_scanned_pages():
    calls: list[tuple[int, int]] = []
    pdf = _pdf_from_images(_scanned_page_image(), _scanned_page_image())

    extract_document(pdf, filename="scan.pdf", progress=lambda done, total: calls.append((done, total)))

    assert calls[-1] == (2, 2)


def test_scanned_pdf_without_ocr_is_flagged_not_silently_empty(monkeypatch):
    import app.services.extraction as extraction_module

    monkeypatch.setattr(extraction_module, "ocr_available", lambda: False)
    result = extract_document(_pdf_from_images(_scanned_page_image()), filename="scan.pdf")

    assert result.pages[0].extraction_method == "ocr_unavailable"
    assert result.pages[0].low_text_warning
    assert result.notes


def test_docx_reads_content_controls_and_text_boxes():
    import copy

    import docx
    from docx.oxml import parse_xml

    document = docx.Document()
    document.add_paragraph("1. SCOPE")
    body = document.element.body
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    body.insert(
        len(body) - 1,
        parse_xml(
            f"<w:sdt {ns}><w:sdtContent><w:p><w:r><w:t>2. TERM of three years</w:t></w:r></w:p>"
            "</w:sdtContent></w:sdt>"
        ),
    )
    textbox_host = document.add_paragraph("3. PRICING")
    textbox = parse_xml(
        f"<w:r {ns}><w:pict><w:txbxContent><w:p><w:r><w:t>Rates are firm</w:t></w:r></w:p>"
        "</w:txbxContent></w:pict></w:r>"
    )
    textbox_host._p.append(copy.deepcopy(textbox))
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_document(buffer.getvalue(), filename="contract.docx").pages[0].text

    assert "1. SCOPE" in text
    assert "2. TERM of three years" in text
    assert "Rates are firm" in text


def test_rtf_is_read_with_its_page_breaks():
    rtf = rb"{\rtf1\ansi {\b 1. SCOPE}\par The Contractor shall test wells.\page 2. TERM\par Three years.}"

    result = extract_document(rtf, filename="contract.rtf")

    assert result.page_count == 2
    assert "The Contractor shall test wells." in result.pages[0].text
    assert "Three years." in result.pages[1].text


def test_odt_is_read_without_libreoffice():
    import zipfile

    content = (
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"'
        ' xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">'
        "<office:body><office:text>"
        "<text:h>1. SCOPE</text:h><text:p>Surface well<text:s/>testing.</text:p>"
        "<text:soft-page-break/>"
        "<table:table><table:table-row>"
        "<table:table-cell><text:p>Day rate</text:p></table:table-cell>"
        "<table:table-cell><text:p>USD 575</text:p></table:table-cell>"
        "</table:table-row></table:table>"
        "</office:text></office:body></office:document-content>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", content)

    result = extract_document(buffer.getvalue(), filename="contract.odt")

    assert result.page_count == 2
    assert "Surface well testing." in result.pages[0].text
    assert result.pages[1].text == "Day rate | USD 575"


def test_legacy_doc_without_word_or_libreoffice_asks_for_docx(monkeypatch):
    import app.services.extraction as extraction_module

    monkeypatch.setattr(extraction_module, "word_available", lambda: False)
    monkeypatch.setattr(extraction_module, "find_libreoffice", lambda: None)

    with pytest.raises(extraction_module.UnsupportedDocumentError, match="save it as .docx"):
        extract_document(b"\xd0\xcf\x11\xe0 legacy word", filename="contract.doc")


def test_docx_automatic_clause_numbering_is_kept():
    # Word keeps "1." / "2." list numbers as formatting, not text; clause
    # segmentation needs them in the text.
    import docx

    document = docx.Document()
    for clause in ("SCOPE", "TERM", "PAYMENT"):
        document.add_paragraph(clause, style="List Number")
    buffer = io.BytesIO()
    document.save(buffer)

    lines = extract_document(buffer.getvalue(), filename="contract.docx").pages[0].text.splitlines()

    assert lines == ["1. SCOPE", "2. TERM", "3. PAYMENT"]


@requires_ocr
def test_clause_numbers_in_a_hanging_indent_margin_stay_with_their_text():
    """Contracts that print clause numbers in their own margin column: OCR
    reads the column as a separate block, detaching every number."""
    image = Image.new("L", (1700, 1400), color=255)
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 34)
    except OSError:
        font = ImageFont.load_default(size=34)
    rows = [
        ("5.2", ["Any taxes, duties, fees and levies payable in", "Pakistan are for the Contractor."]),
        ("5.3", ["The Contractor shall pay all taxes on its", "income under the laws of Pakistan."]),
    ]
    y = 150
    for number, body in rows:
        draw.text((120, y), number, fill=0, font=font)
        for line in body:
            draw.text((330, y), line, fill=0, font=font)
            y += 50
        y += 60

    text = extract_document(_pdf_from_images(image), filename="scan.pdf").pages[0].text

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    assert any(line.startswith("5.2 Any taxes") for line in lines), lines
    assert any(line.startswith("5.3 The Contractor") for line in lines), lines
